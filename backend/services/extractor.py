from __future__ import annotations

from dataclasses import dataclass

import fitz
from pptx import Presentation

from backend.utils.ocr import ocr_pdf


@dataclass
class ExtractionResult:
    text: str
    used_ocr: bool = False


def extract_pdf_text(file_path: str) -> ExtractionResult:
    pages: list[str] = []
    with fitz.open(file_path) as doc:
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text")
            if text and text.strip():
                pages.append(f"[Page {i}]\n{text.strip()}")

    joined = "\n\n".join(pages).strip()
    if joined:
        return ExtractionResult(text=joined, used_ocr=False)

    ocr_text = ocr_pdf(file_path).strip()
    if not ocr_text:
        raise ValueError("Unable to extract text from PDF, including OCR fallback")
    return ExtractionResult(text=ocr_text, used_ocr=True)


def extract_pptx_text(file_path: str) -> ExtractionResult:
    presentation = Presentation(file_path)
    blocks: list[str] = []

    for idx, slide in enumerate(presentation.slides, start=1):
        title = ""
        if slide.shapes.title and slide.shapes.title.text:
            title = slide.shapes.title.text.strip()

        points: list[str] = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                line = paragraph.text.strip()
                if not line:
                    continue
                prefix = "- " if paragraph.level > 0 else ""
                points.append(f"{prefix}{line}")

        section_lines = [f"[Slide {idx}]"]
        if title:
            section_lines.append(f"Title: {title}")
        if points:
            section_lines.append("Points:")
            section_lines.extend(points)

        blocks.append("\n".join(section_lines))

    text = "\n\n".join(blocks).strip()
    if not text:
        raise ValueError("No text content found in PPTX")

    return ExtractionResult(text=text, used_ocr=False)


def extract_text(file_path: str, extension: str) -> ExtractionResult:
    extension = extension.lower().lstrip(".")
    if extension == "pdf":
        return extract_pdf_text(file_path)
    if extension == "pptx":
        return extract_pptx_text(file_path)
    raise ValueError("Unsupported file type")
